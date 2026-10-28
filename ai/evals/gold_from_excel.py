"""정답셋 엑셀(사람이 단 라벨) → structure_gold.json + gold/images/ 복사.

엑셀 '입력' 시트: 한 줄 = 옷 한 벌, 같은 파일명 줄이 여러 개면 한 사진의 여러 벌. 값은 '코드목록' 시트의 이름(앱 공통코드 이름)이라
TrendCodes.java 초안(SEED)의 이름으로 코드값을 찾는다(괄호 앞 이름만 같아도 된다). 코드에 없는 이름이 있으면 멈추고 목록을 보여 준다.
작성 예시 줄(Q열에 표시)과 파일명이 홍길동_로 시작하는 줄은 뺀다. 파일명은 확장자·대소문자가 달라도 이미지 폴더에서 찾는다.

사용(시스템 파이썬 — openpyxl 필요, 가상환경에는 없다):
  py ai/evals/gold_from_excel.py "<엑셀 경로>" "<이미지 폴더>"
기존 structure_gold.json과 gold/images/를 통째로 바꾼다.
"""
import json
import re
import shutil
import sys
import unicodedata
from collections import OrderedDict
from datetime import date
from pathlib import Path

import openpyxl

EVALS = Path(__file__).resolve().parent
GOLD = EVALS / "structure_gold.json"
IMAGES = EVALS / "gold" / "images"
TREND_CODES = next((EVALS.parents[1] / "backend").glob("src/main/java/**/TrendCodes.java"))
COLUMNS = {"category": ("카테고리", "CATEGORY"), "type": ("의류 종류", "ITEM_TYPE"), "materials": ("소재", "MATERIAL"),
           "colors": ("컬러", "COLOR"), "silhouettes": ("실루엣", "SILHOUETTE"), "genders": ("대상", "GENDER"),
           "seasons": ("시즌", "SEASON")}


def nfc(s) -> str:
    return unicodedata.normalize("NFC", str(s)).strip()


def image_ext(path: Path) -> str:
    head = path.read_bytes()[:12]
    if head[:3] == b"\xff\xd8\xff":
        return "jpg"
    if head[:4] == b"\x89PNG":
        return "png"
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return "webp"
    if head[:4] == b"GIF8":
        return "gif"
    return path.suffix.lower().lstrip(".")


# 소재 정답은 사진 질감 기준이다(2026-10-28 사용자 결정 — AI가 사진에서 보는 소재로 적는다). 엑셀 소재는 상품 페이지의 섬유 성분이라
# (청바지 = 면), 사람이 단 종류·아이템 이름에서 질감 소재를 더한다. 섬유 성분은 그대로 두고 채점은 이미지 허용 소재만 본다.
TEXTURE_BY_TYPE = {"BOTTOM_DENIM": "DENIM", "OUTER_DENIM_JACKET": "DENIM", "OUTER_FLEECE": "FLEECE",
                   "TOP_KNIT": "KNIT", "TOP_CARDIGAN": "KNIT"}
TEXTURE_BY_WORD = [(("데님", "denim"), "DENIM"), (("플리스", "후리스", "fleece"), "FLEECE"), (("니트", "knit"), "KNIT"),
                   (("레더", "가죽", "leather"), "LEATHER"), (("스웨이드", "suede"), "SUEDE"),
                   (("시어링", "shearing", "무스탕"), "SHEARLING"), (("코듀로이", "골덴", "corduroy"), "CORDUROY")]


def texture_materials(garment: dict) -> list[str]:
    found = [TEXTURE_BY_TYPE[garment["type"]]] if garment["type"] in TEXTURE_BY_TYPE else []
    name = garment["note"].lower()
    found += [code for words, code in TEXTURE_BY_WORD if any(w in name for w in words)]
    return [m for m in dict.fromkeys(found) if m not in garment["materials"]]


def code_names() -> dict[str, dict[str, str]]:
    """그룹 → {이름: 코드값}. 괄호 앞 이름도 키로 둔다("데님(인디고·청 워싱)" ← "데님")."""
    java = TREND_CODES.read_text(encoding="utf-8")
    seed = java.index("SEED")
    out = {}
    for _, group in COLUMNS.values():
        i = java.index(group + ", List.of(", seed)
        j = re.compile(r"\}\)[,)]").search(java, i).end()
        names = {}
        for value, name in re.findall(r'new String\[\]\{"([^"]+)", "([^"]+)"', java[i:j]):
            names[name] = value
            names.setdefault(name.split("(")[0].strip(), value)
        out[group] = names
    return out


def main():
    xlsx, image_dir = Path(sys.argv[1]), Path(sys.argv[2])
    codes = code_names()
    ws = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)["입력"]
    header = [nfc(h) if h else "" for h in next(ws.iter_rows(max_row=1, values_only=True))]
    col = {h: i for i, h in enumerate(header) if h}
    files = {nfc(f.name).lower(): f for f in image_dir.iterdir() if f.is_file()}
    stems = {nfc(f.stem).lower(): f for f in image_dir.iterdir() if f.is_file()}

    def values(row, label):
        heads = [h for h in col if h == label or re.fullmatch(re.escape(label) + r"\d", h)]
        return [nfc(row[col[h]]) for h in sorted(heads) if row[col[h]] not in (None, "")]

    unknown, missing, docs = set(), [], OrderedDict()
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[0]
        if not name or (len(row) > 16 and row[16]) or nfc(name).startswith("홍길동_"):
            continue
        name = nfc(name)
        image = files.get(name.lower()) or stems.get(Path(name).stem.lower()) or stems.get(name.lower())
        if image is None:
            missing.append(name)
            continue
        garment = {}
        for key, (label, group) in COLUMNS.items():
            found = []
            for v in values(row, label):
                if v not in codes[group]:
                    unknown.add(f"{label}: {v}")
                    continue
                found.append(codes[group][v])
            garment[key] = found
        doc = docs.setdefault(image.name, {"image": image, "source_type": nfc(row[col["자료유형"]] or ""),
                                           "source": nfc(row[col["출처링크"]] or ""), "seasons": [], "garments": []})
        doc["seasons"] += [s for s in garment.pop("seasons") if s not in doc["seasons"]]
        garment["category"] = garment["category"][0] if garment["category"] else None
        garment["type"] = garment["type"][0] if garment["type"] else None
        garment["note"] = nfc(row[col["아이템 이름"]] or "")
        if row[col["비고"]]:
            garment["remark"] = nfc(row[col["비고"]])
        garment["materials"] += texture_materials(garment)
        doc["garments"].append(garment)
    if unknown or missing:
        raise SystemExit("변환 중단 — 코드에 없는 값: " + ", ".join(sorted(unknown)) + " / 이미지 없음: " + ", ".join(missing))

    shutil.rmtree(IMAGES, ignore_errors=True)
    IMAGES.mkdir(parents=True)
    items = []
    for n, (file_name, d) in enumerate(docs.items(), 1):
        # 한글·공백·특수문자 파일명을 피해 번호로 복사한다. 확장자는 내용으로 정한다(.jpg인데 실제 webp인 파일이 있다)
        target = f"G{n:03d}.{image_ext(d['image'])}"
        shutil.copy2(d["image"], IMAGES / target)
        items.append({"id": f"G{n:03d}", "image": f"images/{target}", "file": file_name, "status": "reviewed",
                      "source_type": d["source_type"], "source": d["source"], "seasons": d["seasons"], "garments": d["garments"]})
    GOLD.write_text(json.dumps({"_doc": [
        f"P-03 정형화 정답셋(이미지) — {date.today()} 사람이 엑셀로 단 라벨({xlsx.name})을 gold_from_excel.py로 옮겼다. 엑셀을 고치면 다시 돌린다.",
        "이미지 파일은 gold/images/(커밋 제외 — 매체·쇼핑몰 사진). file은 엑셀의 원래 파일명, source는 출처 링크.",
        "garments: 사진의 옷 한 벌씩, 값은 공통코드 코드값. 빈 배열은 '없음'. note = 엑셀 아이템 이름, remark = 엑셀 비고.",
        "materials는 엑셀의 섬유 성분(면·폴리에스터 등) + 종류·아이템 이름에서 더한 사진 질감 소재(데님·플리스·니트·레더 등 — 사용자 결정: 소재 정답은 사진 질감 기준). 채점은 속성 IMAGE 소재만 본다(structure_eval.py).",
        "seasons는 자료 단위(채점하지 않음). source_type 본문+이미지 자료도 채점은 이미지만으로 한다(본문은 정답셋에 없다).",
    ], "items": items}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"사진 {len(items)}장 · 옷 {sum(len(i['garments']) for i in items)}벌 → {GOLD}")


if __name__ == "__main__":
    main()
